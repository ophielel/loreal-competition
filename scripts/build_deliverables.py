"""Generate the competition deck using the supplied official template."""
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'deliverables'
FONT = 'Noto Sans SC'
INK = '253B35'
TEAL = '17685B'
MUTED = '65756F'


def clear(slide):
    for shape in list(slide.shapes):
        shape._element.getparent().remove(shape._element)


def text(slide, x, y, w, h, value, size=18, color=INK, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(9)
    return shape


def rect(slide, x, y, w, h, fill='F1F6F3'):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    shape.adjustments[0] = .05
    return shape


def title(slide, number, heading, sub=''):
    text(slide,.75,.42,11.8,.28,f'知微  /  BEAUTY CARE COPILOT                                      {number:02d}',10,MUTED)
    text(slide,.75,.85,11.9,.64,heading,28,INK,True)
    if sub:
        text(slide,.78,1.52,11.6,.42,sub,12,MUTED)


def foot(slide, value):
    text(slide,.8,6.37,11.8,.3,value,9,MUTED)


def notes(slide, value):
    slide.notes_slide.notes_text_frame.text = value


def main():
    OUT.mkdir(exist_ok=True)
    p = Presentation(next(ROOT.glob('*PPT*.pptx')))
    cover, content, ending = list(p.slides)
    clear(cover); clear(content); clear(ending)
    # Keep the official master/layout artwork; only replace the provided text areas.
    text(cover,1.25,2.05,11,1.1,'知微',58,'FFFFFF',True)
    text(cover,1.3,3.2,10.6,.7,'消费者服务工作台',32,'FFFFFF',True)
    text(cover,1.32,4.22,10.6,.65,'从看见数据，到理解处境。',23,'FFFFFF')
    text(cover,1.32,5.15,10.7,.8,'赛题一：数据共情者 — 消费者的 AI 管家\n队伍名称：待填写',15,'FFFFFF')
    notes(cover,'参赛作品，不是欧莱雅官方产品。队伍名称按用户要求保留待填写。使用官方提供的 PPT 模板。')
    slides = [content]
    for _ in range(10):
        s = p.slides.add_slide(content.slide_layout)
        clear(s)
        slides.append(s)
    # The original closing slide is moved to the end after all new content slides.
    closing_id = p.slides._sldIdLst[2]
    p.slides._sldIdLst.remove(closing_id)
    p.slides._sldIdLst.append(closing_id)

    s=slides[0]; title(s,1,'问题不止是怎么回复，而是有没有看清。','作品背景 / 面向人工客服在多源信息之间的真实操作负担')
    columns=[('看不清','消息、订单和五类工单分散。','消费者说“上次已经讲过”，客服仍要重新找证据。'),('跟不上','一段对话里，需求和情绪会变化。','单条机械回复看不到催促、重复进线自述与升级线索。'),('判不准','事实、推测和承诺容易混在一起。','售后状态晚于聊天记录时，事后结果可能被误当成当时已知。')]
    for i,(h,a,b) in enumerate(columns):
        x=.8+i*4.05; rect(s,x,2.2,3.8,3.55)
        text(s,x+.25,2.48,3.3,.6,h,27,TEAL,True)
        text(s,x+.25,3.28,3.3,.85,a,18,INK,True)
        text(s,x+.25,4.35,3.3,1.0,b,15,MUTED)
    foot(s,'设计目标：千牛工作台右侧辅助区。依据：官方工作台 DOCX 与赛题说明。')

    s=slides[1]; title(s,2,'从官方数据出发，先尊重它的边界。','作品使用场景 / 全量接入官方虚构业务数据')
    for i,(num,lab) in enumerate([('138','会话'),('998','聊天消息'),('113','订单'),('80','售后工单')]):
        x=.8+i*3.05;rect(s,x,2.15,2.8,1.35);text(s,x+.2,2.3,2.4,.65,num,34,TEAL,True);text(s,x+.22,3,2.3,.3,lab,13,MUTED)
    text(s,.85,3.95,5.7,1.65,'10 类主场景，5 类工单\n会话 ID 关联；来源精确到表格行号\n订单号按字符串保留，避免长数字失真',18)
    rect(s,7.05,3.92,5.4,1.98,'FBF1E4');text(s,7.3,4.1,4.9,.45,'一个影响正确性的发现',18,'88612B',True)
    text(s,7.3,4.7,4.9,1.0,'全部工单晚于对应聊天结束。必须区分历史回放与归档快照。',18,'88612B')
    foot(s,'来源：官方 XLSX。所有内容均为 MOCK；缺失原图、状态变更日志与独立情绪/风险金标。')

    def demo(slide, number, heading, image, points, caption):
        title(slide,number,heading)
        slide.shapes.add_picture(str(OUT/image), Inches(.8), Inches(1.68), height=Inches(4.5))
        for i,(head,body) in enumerate(points):
            y=1.94+i*1.35
            text(slide,8.12,y,4.3,.45,head,18,TEAL,True)
            text(slide,8.12,y+.53,4.3,.78,body,14,MUTED)
        foot(slide,caption)

    demo(slides[2],3,'把辅助能力放回客服最需要的位置。','01-workspace.png',[
        ('左侧：接待分流','按会话检索，筛选高关注内容。'),('中部：保持沟通','原始消息、订单和可编辑回复。'),('右侧：即时辅助','主诉、情绪、风险与下一步建议。')],'完整本地可操作 Demo；截图来自实际 Edge 浏览器，不是设计稿。')
    demo(slides[3],4,'同一段对话，随着证据推进判断。','02-replay.png',[
        ('可控制的历史回放','逐条消息推进，不读取未来回应。'),('情绪只是线索','明确提示启发式判断，避免伪造置信度。'),('风险持续关注','末条“谢谢”不会抹去先前的不适风险。')],'示例 S00010：使用不适场景。原始客服医疗相关话术不代表本系统认可。')
    demo(slides[4],5,'当时知道什么，后来发生什么。','03-journey.png',[
        ('历史回放','按消息时间截断上下文。'),('归档快照','明确标注事后工单与最终记录。'),('三源可追溯','订单、消息、工单组成服务轨迹。')],'无状态日志时不编造历史状态；脱敏昵称不能证明跨会话是同一消费者。')
    demo(slides[5],6,'从一条建议，到一次负责的跟进。','05-tasks.png',[
        ('编辑回复','建议可采用、可修改，模拟发送。'),('人工确认任务','客服检查标题与说明后再保存。'),('本地闭环','SQLite 持久化、相同任务去重、标记完成。')],'未连接真实千牛、物流或支付系统；任务动作不等同于真实退款/补发。')

    s=slides[6]; title(s,7,'Agent 设计：证据在前，行动由人确认。','技术实现 / 受控的分步工作流')
    workflow=[('01  数据关联','聊天 / 订单 / 工单\n会话 ID + 来源行号'),('02  时点门控','历史回放 / 归档快照\n隔离未来信息'),('03  风险证据','semantic_state + evidence\n否定/假设不升级'),('04  默认 Hybrid','Qwen 语义 + 业务规则\n界面内存配置'),('05  安全校验','JSON + 引用 + 回复承诺\n失败时安全回退'),('06  人工闭环','可编辑建议 / 确认任务\n无外部自动执行')]
    for i,(head,body) in enumerate(workflow):
        x=.8+(i%3)*4.06;y=2.12+(i//3)*1.85
        rect(s,x,y,3.8,1.56);text(s,x+.2,y+.18,3.4,.38,head,18,TEAL,True);text(s,x+.2,y+.72,3.4,.7,body,14,MUTED)
    foot(s,'技术栈：Python 标准库 HTTP + 原生 HTML/CSS/JS + SQLite；Qwen 使用兼容 API。详见 AGENT_DESIGN.md。')

    s=slides[7]; title(s,8,'默认 Hybrid 理解语义，证据规则守住边界。','技术实现 / Qwen + Hybrid 默认流程与安全回退')
    for i,(head,body) in enumerate([
        ('输入最小化','仅发送当前可见会话与必要订单/工单字段，不发送整个工作簿。'),
        ('输出可复核','强制 JSON，检查字段、枚举与证据 ID；有效 ID 仍不等于语义充分支持。'),
        ('低成本运行','上下文哈希缓存，最多 256 条；命中不重复请求。只显示实际返回的 tokens。'),
        ('失败可理解','未配置、超时或结果不合法时显示“已安全回退”；密钥仅存服务内存。')]):
        x=.8+(i%2)*6.1;y=2.13+(i//2)*1.85
        rect(s,x,y,5.82,1.62);text(s,x+.22,y+.2,5.3,.42,head,19,TEAL,True);text(s,x+.22,y+.76,5.3,.71,body,15,MUTED)
    foot(s,'已保存 276 个真实 Qwen 评测场景；冲刺版使用保存输出离线重放，仍无真实业务提升或生产泛化结论。')

    release=json.loads((ROOT/'data/release_summary.json').read_text(encoding='utf-8'))
    dev=release['official_dev_metrics']; challenge=release['challenge_metrics']
    s=slides[8]; title(s,9,'先报告可核实的结果，再谈业务价值。','评估 / 官方同源开发集 + 50 条独立人工边界集')
    for i,(num,lab) in enumerate([(f"{dev['first_message']['hybrid']['accuracy']:.1%}",'Hybrid 首条意图'),(f"{dev['full_conversation']['hybrid']['accuracy']:.1%}",'Hybrid 完整会话'),(f"{challenge['unsafe_commitment_recall']:.0%}",'危险承诺拦截召回')]):
        x=.8+i*4.05;rect(s,x,2.12,3.8,1.48);text(s,x+.23,2.25,3.3,.7,num,35,TEAL,True);text(s,x+.23,3.04,3.3,.36,lab,15,MUTED)
    text(s,.85,3.93,5.65,1.96,f"官方 138 会话：Rules {dev['first_message']['rules']['accuracy']:.1%}\nQwen {dev['first_message']['qwen']['accuracy']:.1%} / Hybrid {dev['first_message']['hybrid']['accuracy']:.1%}\n保存的真实 Qwen 输出按当前决策重放\n标签不进入在线推断",17)
    text(s,7.1,3.93,5.25,2.0,f"Challenge Set（50 条，Rules + Safety Gate）\nIntent {challenge['intent_accuracy']:.1%} · Risk Recall {challenge['risk_recall']:.0%}\nEvidence Support {challenge['evidence_support_rate']:.0%}\n小型边界回归，不代表生产泛化。",16,MUTED)
    foot(s,'统一指标来源：data/release_summary.json；详见 docs/EVALUATION.md 与 CHALLENGE_EVALUATION.md。')

    s=slides[9]; title(s,10,'能跑、能核实，也能交接。','项目交付 / 在真实浏览器验证完整操作链路')
    rect(s,.8,2.1,5.7,3.75);text(s,1.05,2.36,5.1,.55,'已交付',23,TEAL,True)
    text(s,1.05,3.06,5.1,2.5,'可运行的三栏工作台\n源码与原始/导入数据\nAgent 设计与运行指南\n官方模板 PPT + 实际操作录屏\n可重跑规则评估与测试',18)
    rect(s,6.8,2.1,5.7,3.75);text(s,7.05,2.36,5.1,.55,'验证证据',23,TEAL,True)
    text(s,7.05,3.06,5.1,2.5,f"{release['python_test_count']} 项 Python 单元 / API 测试\n{release['browser_check_count']} 项真实 Edge 浏览器检查\n320～1440 px 无页面横向溢出\n控制台错误：0\n默认 Hybrid；失败时安全回退",18)
    foot(s,'验证针对本地 Demo；不能替代真实模型、真实千牛接入、安全审计或生产负载测试。')

    s=slides[10]; title(s,11,'把“可能有用”，推进到“证明确实有用”。','未来完善方向 / 价值假设与验证路线')
    road=[('第一步 · 补齐证据','业务专家独立标注\n真实 Qwen 对照评估\n检验证据支持率与误承诺率'),('第二步 · 验证效率','客服交叉顺序试用\n记录找证据耗时与重复提问\n测量真实 tokens 和端到端延迟'),('第三步 · 审慎接入','经授权接入千牛插件\n权限、审计、任务队列\n对高风险动作保留人工审批')]
    for i,(head,body) in enumerate(road):
        x=.8+i*4.05;rect(s,x,2.2,3.8,3.32);text(s,x+.23,2.46,3.33,.75,head,20,TEAL,True);text(s,x+.23,3.46,3.33,1.7,body,17,MUTED)
    foot(s,'目标：减少信息查找、重复询问和无依据承诺。所有业务改善目前为待验证假设。')

    text(ending,1.25,2.08,10.9,.95,'让每一次对话，',40,'FFFFFF',True)
    text(ending,1.25,3.1,10.9,.95,'都被认真理解。',40,'FFFFFF',True)
    text(ending,1.3,4.52,10.7,.9,'知微 · 消费者服务工作台\n队伍名称：待填写',20,'FFFFFF')
    notes(ending,'本地启动：双击 start.cmd 或 python server.py。演示材料位于 deliverables。正式提交前填写队伍名称并核对账号页面限制。')
    for i,slide in enumerate(slides):
        notes(slide,f'内容页 {i+1}。来源：天池赛题 532503、根目录官方 XLSX/DOCX。所有场景为 MOCK 数据。请结合 docs/DEMO_SCRIPT.md 讲解。')
    path=OUT/'知微_参赛展示.pptx'
    p.save(path)
    print(f'Deck created: {len(p.slides)} slides')


if __name__=='__main__':
    main()
